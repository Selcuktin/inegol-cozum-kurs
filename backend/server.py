from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Header
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import logging
import httpx
import requests
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional, Dict
import uuid
from datetime import datetime, timezone, timedelta

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
# MongoDB connection with local fallback
mongo_url = os.environ.get('MONGO_URL', 'mongodb://localhost:27017')
db_name = os.environ.get('DB_NAME', 'test_database')
try:
    import pymongo
    test_client = pymongo.MongoClient(mongo_url, serverSelectionTimeoutMS=1000)
    test_client.server_info()
    client = AsyncIOMotorClient(mongo_url)
    db = client[db_name]
except Exception:
    from mongomock_motor import AsyncMongoMockClient
    client = AsyncMongoMockClient()
    db = client[db_name]

# ----- Roles & Allowlist -----
ROLE_ADMIN = "admin"
ROLE_ACCOUNTING = "muhasebe"
ROLE_TEACHER = "ogretmen"
ROLE_LABELS = {ROLE_ADMIN: "Yönetici", ROLE_ACCOUNTING: "Muhasebe", ROLE_TEACHER: "Öğretmen"}

# ----- Corporate Info (for PDF/Excel) -----
CORP = {
    "name": "ÖZEL İNEGÖL ÇÖZÜM KURS MERKEZİ",
    "owner": "Nevin ÇAKMAK",
    "address": "Kemalpaşa Mah. Ilıca Sokak NO:31/33A İnegöl / BURSA",
    "phones": ["0224 715 00 02", "0553 715 00 02"],
    "tax_no": "3171055818",
    "tax_office": "İnegöl V.D.",
    "slogan": "Problem Yok · Çözüm Var!",
}
LOCAL_LOGO = ROOT_DIR.parent / "frontend" / "public" / "assets" / "logo.jpeg"
LOGO_PATH = LOCAL_LOGO if LOCAL_LOGO.exists() else Path("/app/frontend/public/assets/logo.jpeg")

# ----- Object Storage with Local Fallback -----
LOCAL_STORAGE_DIR = ROOT_DIR / "local_storage"
LOCAL_STORAGE_DIR.mkdir(exist_ok=True)
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "inegol-cozum-kurs"
_storage_key: Optional[str] = None

def init_storage(force: bool = False) -> str:
    global _storage_key
    if _storage_key and not force:
        return _storage_key
    try:
        r = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=5)
        r.raise_for_status()
        _storage_key = r.json()["storage_key"]
        return _storage_key
    except Exception:
        return "local_storage_mode"

def put_object(path: str, data: bytes, content_type: str) -> dict:
    try:
        key = init_storage()
        if key != "local_storage_mode":
            r = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=10)
            if r.status_code == 404:
                key = init_storage(force=True)
                r = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=10)
            r.raise_for_status()
            return r.json()
    except Exception:
        pass
    local_file = LOCAL_STORAGE_DIR / Path(path).name
    local_file.write_bytes(data)
    return {"path": path, "stored": "local"}

def get_object(path: str) -> tuple[bytes, str]:
    try:
        key = init_storage()
        if key != "local_storage_mode":
            r = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=10)
            if r.status_code == 404:
                key = init_storage(force=True)
                r = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=10)
            r.raise_for_status()
            return r.content, r.headers.get("Content-Type", "application/octet-stream")
    except Exception:
        pass
    local_file = LOCAL_STORAGE_DIR / Path(path).name
    if local_file.exists():
        return local_file.read_bytes(), "application/octet-stream"
    raise HTTPException(status_code=404, detail="Dosya bulunamadı")

EMAIL_ROLES = {
    "kayrasemihseymen@gmail.com": ROLE_ADMIN,
    "inegolcozumkurs@gmail.com": ROLE_ADMIN,
    "info@inegolcozumkurs.com": ROLE_ACCOUNTING,
}
ALLOWED_EMAILS = set(EMAIL_ROLES.keys())

# Permission matrix: which roles can access which action
PERMISSIONS = {
    "student.create": {ROLE_ADMIN},
    "payment.create": {ROLE_ADMIN, ROLE_ACCOUNTING},
    "expense.create": {ROLE_ADMIN, ROLE_ACCOUNTING},
    "class.create":   {ROLE_ADMIN},
    "audit.view":     {ROLE_ADMIN},
    "user.manage":    {ROLE_ADMIN},
}

app = FastAPI()
api_router = APIRouter(prefix="/api")


# Define Models
class StatusCheck(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    client_name: str
    timestamp: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))

class StatusCheckCreate(BaseModel):
    client_name: str

class Student(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    student_name: str
    student_phone: str = ""
    school: str = ""
    student_class: str = ""
    student_tc: str = ""
    student_address: str = ""
    parent_name: str
    parent_phone: str = ""
    parent_tc: str = ""
    parent_address: str = ""
    # Anne bilgileri
    mother_tc: str = ""
    mother_name: str = ""
    mother_job: str = ""
    mother_phone: str = ""
    mother_work_phone: str = ""
    mother_address: str = ""
    mother_work_address: str = ""
    mother_email: str = ""
    # Baba (ek alanlar - parent_* mevcut baba bilgisi olarak kullanılır)
    father_job: str = ""
    father_work_phone: str = ""
    father_work_address: str = ""
    father_email: str = ""
    # Vasi
    guardian_tc: str = ""
    guardian_name: str = ""
    guardian_job: str = ""
    guardian_phone: str = ""
    guardian_home_phone: str = ""
    guardian_work_phone: str = ""
    guardian_home_address: str = ""
    guardian_work_address: str = ""
    guardian_email: str = ""
    # Ücret
    program: str = ""
    total_fee: float = 0
    extra_fees: Dict[str, float] = Field(default_factory=lambda: {"kaynak": 0, "dijital": 0, "dnm": 0, "kutuphane": 0, "ozel_ders": 0})
    discount_type: str = ""
    discount_amount: float = 0
    discount_note: str = ""
    paid_amount: float = 0
    installment_count: int = 1
    due_date: str = ""
    photo: str = ""
    # İmzalı sözleşme
    signed_contract_path: str = ""
    signed_contract_filename: str = ""
    signed_contract_uploaded_at: str = ""
    signed_contract_content_type: str = ""
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class PaymentCreate(BaseModel):
    student_id: str
    amount: float
    method: str
    bank: str = ""
    payment_date: str
    note: str = ""

class ExpenseCreate(BaseModel):
    title: str
    amount: float
    expense_date: str
    category: str = "Genel"

class ClassCreate(BaseModel):
    name: str
    level: str = ""
    room: str = ""
    capacity: int = 20
    teacher: str = ""

class User(BaseModel):
    user_id: str
    email: str
    name: str
    picture: str = ""
    role: str = ROLE_TEACHER
    role_label: str = ""

class RoleUpdate(BaseModel):
    user_id: str
    role: str

def clean(doc):
    doc.pop("_id", None)
    return doc

async def seed_data():
    if await db.students.count_documents({}) == 0:
        await db.students.insert_many([
            {"id": "student-001", "student_name": "Eylül Yılmaz", "student_phone": "0532 440 21 18", "school": "İnegöl Anadolu Lisesi", "student_class": "11-A", "student_tc": "", "student_address": "İnegöl / Bursa", "parent_name": "Murat Yılmaz", "parent_phone": "0532 440 21 18", "parent_tc": "", "parent_address": "İnegöl / Bursa", "program": "TYT - AYT", "total_fee": 24000, "paid_amount": 12000, "installment_count": 6, "due_date": "2026-04-15", "photo": "", "created_at": datetime.now(timezone.utc).isoformat()},
            {"id": "student-002", "student_name": "Arda Demir", "student_phone": "0541 309 87 40", "school": "İnegöl Borsa İstanbul", "student_class": "8-B", "student_tc": "", "student_address": "İnegöl / Bursa", "parent_name": "Selin Demir", "parent_phone": "0541 309 87 40", "parent_tc": "", "parent_address": "İnegöl / Bursa", "program": "LGS Hazırlık", "total_fee": 18000, "paid_amount": 6000, "installment_count": 6, "due_date": "2026-03-28", "photo": "", "created_at": datetime.now(timezone.utc).isoformat()}
        ])
    if await db.classes.count_documents({}) == 0:
        await db.classes.insert_many([
            {"id": "class-001", "name": "TYT - AYT", "level": "11. Sınıf", "room": "Derslik 2", "capacity": 24, "teacher": "Ayşe Kaya"},
            {"id": "class-002", "name": "LGS Hazırlık", "level": "8. Sınıf", "room": "Derslik 1", "capacity": 20, "teacher": "Mehmet Çelik"}
        ])


# ---------------- AUTH HELPERS ----------------
async def get_current_user(request: Request) -> User:
    token = request.cookies.get("session_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Yetkisiz")

    session_doc = await db.user_sessions.find_one({"session_token": token}, {"_id": 0})
    if not session_doc:
        raise HTTPException(status_code=401, detail="Oturum bulunamadı")

    expires_at = session_doc.get("expires_at")
    if isinstance(expires_at, str):
        expires_at = datetime.fromisoformat(expires_at)
    if expires_at.tzinfo is None:
        expires_at = expires_at.replace(tzinfo=timezone.utc)
    if expires_at < datetime.now(timezone.utc):
        raise HTTPException(status_code=401, detail="Oturum süresi doldu")

    user_doc = await db.users.find_one({"user_id": session_doc["user_id"]}, {"_id": 0})
    if not user_doc:
        raise HTTPException(status_code=401, detail="Kullanıcı bulunamadı")

    role = user_doc.get("role") or EMAIL_ROLES.get(user_doc.get("email", "").lower(), ROLE_TEACHER)
    user_doc["role"] = role
    user_doc["role_label"] = ROLE_LABELS.get(role, role)
    return User(**user_doc)


def require(user: User, action: str):
    if user.role not in PERMISSIONS.get(action, set()):
        raise HTTPException(status_code=403, detail="Bu işlem için yetkiniz yok")


async def log_action(user: User, action: str, target_type: str = "", target_id: str = "", detail: str = ""):
    await db.audit_logs.insert_one({
        "id": str(uuid.uuid4()),
        "user_id": user.user_id,
        "user_name": user.name,
        "user_email": user.email,
        "user_role": user.role,
        "action": action,
        "target_type": target_type,
        "target_id": target_id,
        "detail": detail,
        "created_at": datetime.now(timezone.utc).isoformat(),
    })


# ---------------- AUTH ROUTES ----------------
@api_router.post("/auth/session")
async def auth_session(request: Request, response: Response, x_session_id: Optional[str] = Header(None)):
    if not x_session_id:
        raise HTTPException(status_code=400, detail="X-Session-ID header gerekli")

    async with httpx.AsyncClient(timeout=15) as http:
        r = await http.get(
            "https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data",
            headers={"X-Session-ID": x_session_id},
        )
    if r.status_code != 200:
        raise HTTPException(status_code=401, detail="Geçersiz session_id")
    data = r.json()

    email = (data.get("email") or "").lower().strip()
    if email not in ALLOWED_EMAILS:
        raise HTTPException(status_code=403, detail=f"{email} adresi bu sisteme erişim iznine sahip değil. Yetkili e-posta ile giriş yapın.")

    name = data.get("name", "")
    picture = data.get("picture", "")
    session_token = data["session_token"]
    assigned_role = EMAIL_ROLES[email]

    existing = await db.users.find_one({"email": email}, {"_id": 0})
    if existing:
        user_id = existing["user_id"]
        current_role = existing.get("role") or assigned_role
        await db.users.update_one({"user_id": user_id}, {"$set": {"name": name, "picture": picture, "role": current_role}})
    else:
        user_id = f"user_{uuid.uuid4().hex[:12]}"
        current_role = assigned_role
        await db.users.insert_one({
            "user_id": user_id,
            "email": email,
            "name": name,
            "picture": picture,
            "role": current_role,
            "created_at": datetime.now(timezone.utc),
        })

    await db.user_sessions.insert_one({
        "user_id": user_id,
        "session_token": session_token,
        "expires_at": datetime.now(timezone.utc) + timedelta(days=7),
        "created_at": datetime.now(timezone.utc),
    })

    fake_user = User(user_id=user_id, email=email, name=name, picture=picture, role=current_role, role_label=ROLE_LABELS.get(current_role, current_role))
    await log_action(fake_user, "auth.login", "user", user_id, f"{name} giriş yaptı")

    response.set_cookie(
        key="session_token", value=session_token,
        httponly=True, secure=True, samesite="none", path="/",
        max_age=7 * 24 * 60 * 60,
    )
    return {"user_id": user_id, "email": email, "name": name, "picture": picture, "role": current_role, "role_label": ROLE_LABELS.get(current_role, current_role)}


@api_router.get("/auth/me")
async def auth_me(request: Request):
    user = await get_current_user(request)
    return user.model_dump()


@api_router.post("/auth/logout")
async def auth_logout(request: Request, response: Response):
    token = request.cookies.get("session_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if token:
        session_doc = await db.user_sessions.find_one({"session_token": token}, {"_id": 0})
        if session_doc:
            user_doc = await db.users.find_one({"user_id": session_doc["user_id"]}, {"_id": 0})
            if user_doc:
                await db.audit_logs.insert_one({
                    "id": str(uuid.uuid4()), "user_id": user_doc["user_id"], "user_name": user_doc.get("name",""),
                    "user_email": user_doc.get("email",""), "user_role": user_doc.get("role",""),
                    "action": "auth.logout", "target_type": "user", "target_id": user_doc["user_id"],
                    "detail": "Çıkış yaptı", "created_at": datetime.now(timezone.utc).isoformat()
                })
        await db.user_sessions.delete_one({"session_token": token})
    response.delete_cookie("session_token", path="/")
    return {"ok": True}


class QuickLoginRequest(BaseModel):
    role: str = ROLE_ADMIN

@api_router.post("/auth/quick-login")
async def quick_login(payload: QuickLoginRequest, request: Request, response: Response):
    role = payload.role if payload.role in ROLE_LABELS else ROLE_ADMIN
    demo_users = {
        ROLE_ADMIN: ("kayrasemihseymen@gmail.com", "Semih Seymen (Yönetici)"),
        ROLE_ACCOUNTING: ("info@inegolcozumkurs.com", "Muhasebe Sorumlusu"),
        ROLE_TEACHER: ("ogretmen@inegolcozumkurs.com", "Ayşe Kaya (Öğretmen)"),
    }
    email, name = demo_users.get(role, ("kayrasemihseymen@gmail.com", "Yönetici"))

    existing = await db.users.find_one({"email": email}, {"_id": 0})
    if existing:
        user_id = existing["user_id"]
        await db.users.update_one({"user_id": user_id}, {"$set": {"name": name, "role": role}})
    else:
        user_id = f"user_{uuid.uuid4().hex[:12]}"
        await db.users.insert_one({
            "user_id": user_id,
            "email": email,
            "name": name,
            "picture": "",
            "role": role,
            "created_at": datetime.now(timezone.utc),
        })

    session_token = f"quick_session_{uuid.uuid4().hex}"
    await db.user_sessions.insert_one({
        "user_id": user_id,
        "session_token": session_token,
        "expires_at": datetime.now(timezone.utc) + timedelta(days=7),
        "created_at": datetime.now(timezone.utc),
    })

    fake_user = User(user_id=user_id, email=email, name=name, picture="", role=role, role_label=ROLE_LABELS.get(role, role))
    await log_action(fake_user, "auth.quick_login", "user", user_id, f"{name} ({role}) hızlı giriş yaptı")

    is_https = request.url.scheme == "https"
    response.set_cookie(
        key="session_token", value=session_token,
        httponly=True, secure=is_https, samesite="none" if is_https else "lax", path="/",
        max_age=7 * 24 * 60 * 60,
    )
    return {
        "user_id": user_id,
        "email": email,
        "name": name,
        "picture": "",
        "role": role,
        "role_label": ROLE_LABELS.get(role, role),
        "session_token": session_token
    }


# ---------------- USERS & AUDIT ----------------
@api_router.get("/users")
async def list_users(request: Request):
    user = await get_current_user(request)
    require(user, "user.manage")
    docs = await db.users.find({}, {"_id": 0}).to_list(500)
    for d in docs:
        d["role_label"] = ROLE_LABELS.get(d.get("role", ROLE_TEACHER), d.get("role", ""))
    return docs


@api_router.put("/users/role")
async def update_role(input: RoleUpdate, request: Request):
    user = await get_current_user(request)
    require(user, "user.manage")
    if input.role not in ROLE_LABELS:
        raise HTTPException(422, "Geçersiz rol")
    if input.user_id == user.user_id and input.role != ROLE_ADMIN:
        raise HTTPException(400, "Kendi yönetici rolünüzü kaldıramazsınız")
    await db.users.update_one({"user_id": input.user_id}, {"$set": {"role": input.role}})
    await log_action(user, "user.role_update", "user", input.user_id, f"Rol {ROLE_LABELS[input.role]} olarak güncellendi")
    return {"ok": True}


@api_router.get("/audit-logs")
async def get_audit_logs(request: Request, limit: int = 200):
    user = await get_current_user(request)
    require(user, "audit.view")
    docs = await db.audit_logs.find({}, {"_id": 0}).sort("created_at", -1).to_list(limit)
    return docs


# ---------------- STUDENT DETAIL ----------------
def compute_installment_plan(student: dict, payments_for_student: list) -> list:
    n = int(student.get("installment_count") or 1)
    total = float(student.get("total_fee") or 0)
    if n < 1: n = 1
    per = round(total / n, 2)
    try:
        start = datetime.fromisoformat(str(student.get("due_date") or "")[:10])
    except Exception:
        start = datetime.now(timezone.utc)
    paid_remaining = float(student.get("paid_amount") or 0)
    plan = []
    for i in range(n):
        month = start.month + i
        year = start.year + (month - 1) // 12
        month = ((month - 1) % 12) + 1
        try:
            due = datetime(year, month, min(start.day, 28))
        except Exception:
            due = start
        applied = min(per, paid_remaining)
        paid_remaining = max(0, paid_remaining - per)
        plan.append({
            "no": i + 1,
            "due_date": due.strftime("%Y-%m-%d"),
            "amount": per,
            "paid": round(applied, 2),
            "remaining": round(per - applied, 2),
            "status": "paid" if applied >= per - 0.5 else ("partial" if applied > 0 else "pending"),
        })
    return plan


@api_router.get("/students/{student_id}/detail")
async def student_detail(student_id: str, request: Request):
    await get_current_user(request)
    student = await db.students.find_one({"id": student_id}, {"_id": 0})
    if not student:
        raise HTTPException(404, "Öğrenci bulunamadı")
    student["remaining"] = max(0, (student.get("total_fee") or 0) - (student.get("paid_amount") or 0))
    payments = await db.payments.find({"student_id": student_id}, {"_id": 0}).sort("payment_date", -1).to_list(500)
    plan = compute_installment_plan(student, payments)
    return {"student": student, "payments": payments, "installment_plan": plan}


# ---------------- REMINDERS ----------------
@api_router.get("/reminders")
async def reminders(request: Request, days: int = 7):
    await get_current_user(request)
    students = await db.students.find({}, {"_id": 0}).to_list(1000)
    payments = await db.payments.find({}, {"_id": 0}).to_list(2000)
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    horizon = now + timedelta(days=days)
    items = []
    for s in students:
        remaining = (s.get("total_fee") or 0) - (s.get("paid_amount") or 0)
        if remaining <= 0:
            continue
        student_payments = [p for p in payments if p.get("student_id") == s.get("id")]
        plan = compute_installment_plan(s, student_payments)
        next_due = None
        for p in plan:
            if p["status"] != "paid":
                next_due = p; break
        if not next_due:
            continue
        try:
            due_dt = datetime.fromisoformat(next_due["due_date"])
        except Exception:
            continue
        status = "overdue" if due_dt < now else ("upcoming" if due_dt <= horizon else "future")
        if status == "future":
            continue
        items.append({
            "student_id": s.get("id"),
            "student_name": s.get("student_name"),
            "parent_name": s.get("parent_name"),
            "parent_phone": s.get("parent_phone") or s.get("student_phone") or "",
            "installment_no": next_due["no"],
            "amount_due": next_due["remaining"] or next_due["amount"],
            "due_date": next_due["due_date"],
            "status": status,
            "days_diff": (due_dt - now).days,
        })
    items.sort(key=lambda x: (x["status"] != "overdue", x["due_date"]))
    return items


# ---------------- EXCEL EXPORT ----------------
def _autofit(ws):
    for col in ws.columns:
        max_len = 8
        letter = col[0].column_letter
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[letter].width = min(48, max_len + 2)


def _header_row(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(vertical="center")


async def _build_workbook(kind: str) -> Workbook:
    wb = Workbook()
    ws0 = wb.active
    wb.remove(ws0)

    async def sheet_students():
        ws = wb.create_sheet("Öğrenciler")
        _header_row(ws, ["Ad Soyad", "Telefon", "TC", "Okul", "Sınıf", "Veli", "Veli Telefon", "Program", "Toplam Ücret", "Ödenen", "Kalan", "Taksit", "İlk Vade"])
        for s in await db.students.find({}, {"_id": 0}).to_list(2000):
            remaining = (s.get("total_fee") or 0) - (s.get("paid_amount") or 0)
            ws.append([s.get("student_name",""), s.get("student_phone",""), s.get("student_tc",""), s.get("school",""), s.get("student_class",""), s.get("parent_name",""), s.get("parent_phone",""), s.get("program",""), s.get("total_fee",0), s.get("paid_amount",0), remaining, s.get("installment_count",1), s.get("due_date","")])
        _autofit(ws)

    async def sheet_payments():
        ws = wb.create_sheet("Tahsilatlar")
        _header_row(ws, ["Tarih", "Öğrenci", "Tutar (₺)", "Yöntem", "Banka", "Açıklama", "Kaydeden"])
        smap = {s["id"]: s for s in await db.students.find({}, {"_id": 0}).to_list(2000)}
        for p in await db.payments.find({}, {"_id": 0}).sort("payment_date", -1).to_list(5000):
            ws.append([p.get("payment_date",""), smap.get(p.get("student_id",""),{}).get("student_name","-"), p.get("amount",0), p.get("method",""), p.get("bank",""), p.get("note",""), p.get("created_by_name","")])
        _autofit(ws)

    async def sheet_expenses():
        ws = wb.create_sheet("Giderler")
        _header_row(ws, ["Tarih", "Başlık", "Kategori", "Tutar (₺)", "Kaydeden"])
        for e in await db.expenses.find({}, {"_id": 0}).sort("expense_date", -1).to_list(5000):
            ws.append([e.get("expense_date",""), e.get("title",""), e.get("category",""), e.get("amount",0), e.get("created_by_name","")])
        _autofit(ws)

    async def sheet_audit():
        ws = wb.create_sheet("Denetim")
        _header_row(ws, ["Tarih", "Kullanıcı", "E-posta", "Rol", "İşlem", "Açıklama"])
        for a in await db.audit_logs.find({}, {"_id": 0}).sort("created_at", -1).to_list(5000):
            ws.append([a.get("created_at",""), a.get("user_name",""), a.get("user_email",""), ROLE_LABELS.get(a.get("user_role",""), a.get("user_role","")), a.get("action",""), a.get("detail","")])
        _autofit(ws)

    mapping = {"students": [sheet_students], "payments": [sheet_payments], "expenses": [sheet_expenses], "audit": [sheet_audit],
               "all": [sheet_students, sheet_payments, sheet_expenses, sheet_audit]}
    for fn in mapping.get(kind, []):
        await fn()
    return wb


@api_router.get("/export/{kind}")
async def export_excel(kind: str, request: Request):
    user = await get_current_user(request)
    if kind == "audit" and user.role != ROLE_ADMIN:
        raise HTTPException(403, "Bu dışa aktarım için yetkiniz yok")
    if kind not in {"students", "payments", "expenses", "audit", "all"}:
        raise HTTPException(404, "Bilinmeyen dışa aktarım türü")
    wb = await _build_workbook(kind)
    buf = io.BytesIO()
    wb.save(buf); buf.seek(0)
    filename = f"inegol-cozum-{kind}-{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
    await log_action(user, "export.excel", "export", kind, f"{kind} Excel dışa aktarıldı")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


# ---------------- PDF MONTHLY REPORT ----------------
def _pdf_corp_header(story, styles):
    header_data = []
    logo_cell = ""
    if LOGO_PATH.exists():
        try:
            logo_cell = RLImage(str(LOGO_PATH), width=26*mm, height=26*mm)
        except Exception:
            logo_cell = ""
    corp_html = (
        f"<b><font size=13 color='#2563EB'>{CORP['name']}</font></b><br/>"
        f"<font size=9 color='#64748b'>{CORP['owner']} · {CORP['slogan']}</font><br/>"
        f"<font size=9>{CORP['address']}</font><br/>"
        f"<font size=9>Tel: {' / '.join(CORP['phones'])} · Vergi No: {CORP['tax_no']} ({CORP['tax_office']})</font>"
    )
    header_data.append([logo_cell, Paragraph(corp_html, styles["body"])])
    t = Table(header_data, colWidths=[30*mm, 145*mm])
    t.setStyle(TableStyle([
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("BOTTOMPADDING", (0,0), (-1,-1), 6),
        ("LINEBELOW", (0,0), (-1,-1), 1.5, colors.HexColor("#0f172a")),
    ]))
    story.append(t)
    story.append(Spacer(1, 6*mm))


@api_router.get("/reports/monthly")
async def monthly_report(request: Request, month: Optional[str] = None):
    user = await get_current_user(request)
    if not month:
        month = datetime.now(timezone.utc).strftime("%Y-%m")

    students = await db.students.find({}, {"_id": 0}).to_list(2000)
    smap = {s["id"]: s for s in students}
    payments = await db.payments.find({}, {"_id": 0}).sort("payment_date", 1).to_list(5000)
    expenses = await db.expenses.find({}, {"_id": 0}).sort("expense_date", 1).to_list(5000)

    period_payments = [p for p in payments if str(p.get("payment_date","")).startswith(month)]
    period_expenses = [e for e in expenses if str(e.get("expense_date","")).startswith(month)]
    collected = sum(p.get("amount",0) for p in period_payments)
    exp_total = sum(e.get("amount",0) for e in period_expenses)
    outstanding = sum(max(0, (s.get("total_fee",0) - s.get("paid_amount",0))) for s in students)

    method_breakdown = {}
    for p in period_payments:
        method_breakdown[p.get("method","Diğer")] = method_breakdown.get(p.get("method","Diğer"), 0) + (p.get("amount") or 0)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15*mm, rightMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm, title=f"Aylık Rapor {month}")
    ss = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle("t", parent=ss["Title"], fontSize=18, textColor=colors.HexColor("#0f172a"), spaceAfter=4),
        "sub":   ParagraphStyle("s", parent=ss["Normal"], fontSize=9, textColor=colors.HexColor("#64748b"), spaceAfter=10),
        "h2":    ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12, textColor=colors.HexColor("#2563EB"), spaceBefore=6, spaceAfter=6),
        "body":  ParagraphStyle("b", parent=ss["Normal"], fontSize=9),
    }
    story = []
    _pdf_corp_header(story, styles)
    story.append(Paragraph(f"AYLIK MUHASEBE RAPORU · {month}", styles["title"]))
    story.append(Paragraph(f"Oluşturan: {user.name} ({ROLE_LABELS.get(user.role, user.role)}) · {datetime.now(timezone.utc).strftime('%d.%m.%Y %H:%M')} UTC", styles["sub"]))

    net = collected - exp_total
    summary_data = [
        ["", "Bu Ay", ""],
        ["Toplam Tahsilat", f"{collected:,.2f} ₺", f"{len(period_payments)} işlem"],
        ["Toplam Gider", f"{exp_total:,.2f} ₺", f"{len(period_expenses)} kayıt"],
        ["Net", f"{net:,.2f} ₺", "Tahsilat − Gider"],
        ["Toplam Bekleyen Alacak", f"{outstanding:,.2f} ₺", f"{len(students)} öğrenci geneli"],
    ]
    t = Table(summary_data, colWidths=[70*mm, 55*mm, 50*mm])
    t.setStyle(TableStyle([
        ("SPAN", (0,0), (-1,0)),
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#2563EB")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("ALIGN", (0,0), (-1,0), "CENTER"),
        ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#cbd5e1")),
        ("BACKGROUND", (0,1), (0,-1), colors.HexColor("#f1f5f9")),
        ("FONTNAME", (0,1), (0,-1), "Helvetica-Bold"),
        ("TEXTCOLOR", (1,3), (1,3), colors.HexColor("#059669" if net >= 0 else "#dc2626")),
        ("FONTSIZE", (0,0), (-1,-1), 9),
        ("BOTTOMPADDING", (0,0), (-1,-1), 6),
        ("TOPPADDING", (0,0), (-1,-1), 6),
    ]))
    story.append(t)
    story.append(Spacer(1, 6*mm))

    if method_breakdown:
        story.append(Paragraph("Ödeme Yöntemi Dağılımı", styles["h2"]))
        mb = [["Yöntem", "Tutar", "Oran"]]
        for m, v in sorted(method_breakdown.items(), key=lambda x: -x[1]):
            pct = (v / collected * 100) if collected else 0
            mb.append([m, f"{v:,.2f} ₺", f"%{pct:.1f}"])
        t = Table(mb, colWidths=[80*mm, 45*mm, 50*mm])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#eff6ff")),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("GRID", (0,0), (-1,-1), 0.4, colors.HexColor("#cbd5e1")),
            ("FONTSIZE", (0,0), (-1,-1), 9),
        ]))
        story.append(t)
        story.append(Spacer(1, 6*mm))

    story.append(Paragraph(f"Tahsilat Detayı ({len(period_payments)} kayıt)", styles["h2"]))
    if period_payments:
        rows = [["Tarih", "Öğrenci", "Yöntem", "Tutar", "Açıklama"]]
        for p in period_payments[:60]:
            rows.append([
                p.get("payment_date",""),
                (smap.get(p.get("student_id",""),{}).get("student_name","-"))[:28],
                p.get("method",""),
                f"{(p.get('amount') or 0):,.2f} ₺",
                (p.get("note","") or "")[:32],
            ])
        t = Table(rows, colWidths=[22*mm, 60*mm, 30*mm, 30*mm, 33*mm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0f172a")),
            ("TEXTCOLOR", (0,0), (-1,0), colors.white),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE", (0,0), (-1,-1), 8),
            ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#e2e8f0")),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f8fafc")]),
        ]))
        story.append(t)
    else:
        story.append(Paragraph("<i>Bu ay için tahsilat kaydı bulunamadı.</i>", styles["body"]))
    story.append(Spacer(1, 6*mm))

    story.append(Paragraph(f"Gider Detayı ({len(period_expenses)} kayıt)", styles["h2"]))
    if period_expenses:
        rows = [["Tarih", "Başlık", "Kategori", "Tutar"]]
        for e in period_expenses[:60]:
            rows.append([e.get("expense_date",""), (e.get("title",""))[:40], e.get("category",""), f"{(e.get('amount') or 0):,.2f} ₺"])
        t = Table(rows, colWidths=[25*mm, 80*mm, 40*mm, 30*mm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#0f172a")),
            ("TEXTCOLOR", (0,0), (-1,0), colors.white),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE", (0,0), (-1,-1), 8),
            ("GRID", (0,0), (-1,-1), 0.3, colors.HexColor("#e2e8f0")),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f8fafc")]),
        ]))
        story.append(t)
    else:
        story.append(Paragraph("<i>Bu ay için gider kaydı bulunamadı.</i>", styles["body"]))

    story.append(Spacer(1, 10*mm))
    story.append(Paragraph(f"<font size=7 color='#94a3b8'>Bu rapor {datetime.now(timezone.utc).strftime('%d.%m.%Y %H:%M')} UTC tarihinde İnegöl Çözüm Kurs yönetim paneli tarafından otomatik oluşturulmuştur.</font>", styles["body"]))

    doc.build(story)
    buf.seek(0)
    await log_action(user, "report.monthly", "report", month, f"{month} aylık PDF rapor indirildi")
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="inegol-cozum-rapor-{month}.pdf"'})


# ---------------- APP ROUTES ----------------
@api_router.get("/")
async def root():
    return {"message": "İnegöl Çözüm Kurs API"}

@api_router.get("/dashboard")
async def dashboard(request: Request):
    await get_current_user(request)
    await seed_data()
    students = await db.students.find({}, {"_id": 0}).to_list(1000)
    payments = await db.payments.find({}, {"_id": 0}).sort("payment_date", -1).to_list(1000)
    expenses = await db.expenses.find({}, {"_id": 0}).sort("expense_date", -1).to_list(1000)
    classes = await db.classes.find({}, {"_id": 0}).to_list(1000)
    for s in students:
        s["remaining"] = max(0, s.get("total_fee", 0) - s.get("paid_amount", 0))
    month_key = datetime.now(timezone.utc).strftime("%Y-%m")
    period_payments = [p for p in payments if str(p.get("payment_date", "")).startswith(month_key)]
    period_expenses = [e for e in expenses if str(e.get("expense_date", "")).startswith(month_key)]
    return {"students": students, "payments": payments, "expenses": expenses, "classes": classes,
            "summary": {"period": month_key, "collected": sum(p.get("amount", 0) for p in period_payments),
                        "expenses": sum(e.get("amount", 0) for e in period_expenses)}}

@api_router.get("/students", response_model=List[Student])
async def get_students(request: Request, q: Optional[str] = None):
    await get_current_user(request)
    await seed_data()
    query = {}
    if q:
        query = {"$or": [{"student_name": {"$regex": q, "$options": "i"}}, {"parent_name": {"$regex": q, "$options": "i"}}, {"student_phone": {"$regex": q, "$options": "i"}}]}
    return await db.students.find(query, {"_id": 0}).to_list(1000)

@api_router.post("/students", response_model=Student)
async def create_student(input: Student, request: Request):
    user = await get_current_user(request)
    require(user, "student.create")
    doc = input.model_dump()
    doc["created_by"] = user.user_id
    doc["created_by_name"] = user.name
    await db.students.insert_one(doc)
    await log_action(user, "student.create", "student", input.id, f"{input.student_name} kaydedildi")
    return input

@api_router.get("/students/{student_id}", response_model=Student)
async def get_student(student_id: str, request: Request):
    await get_current_user(request)
    doc = await db.students.find_one({"id": student_id}, {"_id": 0})
    if not doc: raise HTTPException(404, "Öğrenci bulunamadı")
    return doc

@api_router.patch("/students/{student_id}", response_model=Student)
async def update_student(student_id: str, patch: dict, request: Request):
    user = await get_current_user(request)
    require(user, "student.create")
    existing = await db.students.find_one({"id": student_id}, {"_id": 0})
    if not existing: raise HTTPException(404, "Öğrenci bulunamadı")
    allowed = {k: v for k, v in patch.items() if k in Student.model_fields and k not in {"id", "created_at", "paid_amount"}}
    if allowed:
        await db.students.update_one({"id": student_id}, {"$set": allowed})
    await log_action(user, "student.update", "student", student_id, f"{existing.get('student_name','')} güncellendi ({len(allowed)} alan)")
    doc = await db.students.find_one({"id": student_id}, {"_id": 0})
    return doc

@api_router.post("/students/{student_id}/contract-upload")
async def upload_contract(student_id: str, request: Request):
    user = await get_current_user(request)
    require(user, "student.create")
    existing = await db.students.find_one({"id": student_id}, {"_id": 0})
    if not existing: raise HTTPException(404, "Öğrenci bulunamadı")
    body = await request.body()
    if not body:
        raise HTTPException(400, "Dosya boş")
    if len(body) > 10 * 1024 * 1024:
        raise HTTPException(413, "Dosya 10MB'dan büyük olamaz")
    content_type = request.headers.get("X-Content-Type") or request.headers.get("Content-Type") or "application/octet-stream"
    if content_type.startswith("multipart"):
        content_type = "application/pdf"
    filename = request.headers.get("X-Filename", f"sozlesme-{student_id}.bin")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "bin"
    if ext not in {"pdf", "jpg", "jpeg", "png"}:
        raise HTTPException(400, "Sadece PDF, JPG veya PNG kabul edilir")
    path = f"{APP_NAME}/contracts/{student_id}/{uuid.uuid4()}.{ext}"
    try:
        result = put_object(path, body, content_type)
    except Exception as e:
        logger.exception("Object storage upload failed")
        raise HTTPException(502, f"Sözleşme yüklenemedi: {e}")
    now = datetime.now(timezone.utc).isoformat()
    await db.students.update_one({"id": student_id}, {"$set": {
        "signed_contract_path": result["path"],
        "signed_contract_filename": filename,
        "signed_contract_content_type": content_type,
        "signed_contract_uploaded_at": now,
    }})
    await log_action(user, "contract.upload", "student", student_id, f"{existing.get('student_name','')} sözleşmesi yüklendi ({filename})")
    return {"ok": True, "path": result["path"], "filename": filename, "uploaded_at": now}

@api_router.get("/students/{student_id}/contract-file")
async def get_contract(student_id: str, request: Request):
    await get_current_user(request)
    s = await db.students.find_one({"id": student_id}, {"_id": 0})
    if not s or not s.get("signed_contract_path"):
        raise HTTPException(404, "İmzalı sözleşme yok")
    try:
        data, ct = get_object(s["signed_contract_path"])
    except Exception as e:
        raise HTTPException(502, f"Sözleşme okunamadı: {e}")
    ct = s.get("signed_contract_content_type") or ct
    return Response(content=data, media_type=ct, headers={"Content-Disposition": f'inline; filename="{s.get("signed_contract_filename","sozlesme.pdf")}"'})

@api_router.post("/payments")
async def create_payment(input: PaymentCreate, request: Request):
    user = await get_current_user(request)
    require(user, "payment.create")
    student = await db.students.find_one({"id": input.student_id}, {"_id": 0})
    if not student: raise HTTPException(404, "Öğrenci bulunamadı")
    if input.amount <= 0: raise HTTPException(422, "Tahsilat tutarı sıfırdan büyük olmalıdır")
    remaining = student.get("total_fee", 0) - student.get("paid_amount", 0)
    if input.amount > remaining: raise HTTPException(422, f"Tahsilat kalan bakiyeyi ({remaining:.2f} ₺) aşamaz")
    payment = {"id": str(uuid.uuid4()), **input.model_dump(), "created_at": datetime.now(timezone.utc).isoformat(), "created_by": user.user_id, "created_by_name": user.name}
    await db.payments.insert_one(payment)
    await db.students.update_one({"id": input.student_id}, {"$inc": {"paid_amount": input.amount}})
    await log_action(user, "payment.create", "payment", payment["id"], f"{student.get('student_name','?')} · {input.amount:.0f} ₺ · {input.method}")
    return clean(payment)

@api_router.post("/expenses")
async def create_expense(input: ExpenseCreate, request: Request):
    user = await get_current_user(request)
    require(user, "expense.create")
    expense = {"id": str(uuid.uuid4()), **input.model_dump(), "created_by": user.user_id, "created_by_name": user.name}
    await db.expenses.insert_one(expense)
    await log_action(user, "expense.create", "expense", expense["id"], f"{input.title} · {input.amount:.0f} ₺")
    return clean(expense)

@api_router.post("/classes")
async def create_class(input: ClassCreate, request: Request):
    user = await get_current_user(request)
    require(user, "class.create")
    item = {"id": str(uuid.uuid4()), **input.model_dump(), "created_by": user.user_id, "created_by_name": user.name}
    await db.classes.insert_one(item)
    await log_action(user, "class.create", "class", item["id"], f"{input.name} sınıfı oluşturuldu")
    return clean(item)

@api_router.post("/status", response_model=StatusCheck)
async def create_status_check(input: StatusCheckCreate):
    status_dict = input.model_dump()
    status_obj = StatusCheck(**status_dict)
    doc = status_obj.model_dump()
    doc['timestamp'] = doc['timestamp'].isoformat()
    _ = await db.status_checks.insert_one(doc)
    return status_obj

@api_router.get("/status", response_model=List[StatusCheck])
async def get_status_checks():
    status_checks = await db.status_checks.find({}, {"_id": 0}).to_list(1000)
    for check in status_checks:
        if isinstance(check['timestamp'], str):
            check['timestamp'] = datetime.fromisoformat(check['timestamp'])
    return status_checks

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origin_regex=".*",
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()

@app.on_event("startup")
async def startup_event():
    try:
        await seed_data()
        logger.info("Seed data verified")
    except Exception as e:
        logger.warning(f"Seed data error: {e}")
    try:
        init_storage()
        logger.info("Object storage initialized")
    except Exception as e:
        logger.warning(f"Object storage init deferred: {e}")
