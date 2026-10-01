"""Tests for: monthly PDF report, Excel exports, reminders, student detail drawer, audit."""
import os
import io
import pytest
import requests
from openpyxl import load_workbook
from pymongo import MongoClient

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', 'https://student-portal-track.preview.emergentagent.com').rstrip('/')
TOKEN = "test_session_inegol_123"
USER_ID = "test-user-inegol"
MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")

_mongo = MongoClient(MONGO_URL)
_db = _mongo[DB_NAME]


def set_role(role):
    _db.users.update_one({"user_id": USER_ID}, {"$set": {"role": role}})


@pytest.fixture(scope="module", autouse=True)
def restore_admin():
    set_role("admin")
    yield
    set_role("admin")


@pytest.fixture
def client():
    s = requests.Session()
    s.headers.update({"Authorization": f"Bearer {TOKEN}"})
    return s


# ------- Monthly PDF Report -------
class TestMonthlyReport:
    def test_admin_pdf_ok(self, client):
        set_role("admin")
        r = client.get(f"{BASE_URL}/api/reports/monthly?month=2026-03")
        assert r.status_code == 200
        assert "application/pdf" in r.headers.get("content-type", "")
        assert r.content.startswith(b"%PDF")
        assert len(r.content) > 500

    def test_muhasebe_pdf_ok(self, client):
        set_role("muhasebe")
        r = client.get(f"{BASE_URL}/api/reports/monthly?month=2026-03")
        set_role("admin")
        assert r.status_code == 200
        assert r.content.startswith(b"%PDF")

    def test_ogretmen_pdf_ok(self, client):
        set_role("ogretmen")
        r = client.get(f"{BASE_URL}/api/reports/monthly?month=2026-03")
        set_role("admin")
        assert r.status_code == 200
        assert r.content.startswith(b"%PDF")

    def test_no_month_defaults(self, client):
        r = client.get(f"{BASE_URL}/api/reports/monthly")
        assert r.status_code == 200
        assert r.content.startswith(b"%PDF")


# ------- Excel exports -------
class TestExcelExports:
    def test_export_students(self, client):
        r = client.get(f"{BASE_URL}/api/export/students")
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        assert "Öğrenciler" in wb.sheetnames
        ws = wb["Öğrenciler"]
        # header + N students
        student_count = _db.students.count_documents({})
        assert ws.max_row == student_count + 1

    def test_export_payments(self, client):
        r = client.get(f"{BASE_URL}/api/export/payments")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert "Tahsilatlar" in wb.sheetnames

    def test_export_expenses(self, client):
        r = client.get(f"{BASE_URL}/api/export/expenses")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert "Giderler" in wb.sheetnames

    def test_export_all_has_4_sheets(self, client):
        r = client.get(f"{BASE_URL}/api/export/all")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert set(wb.sheetnames) >= {"Öğrenciler", "Tahsilatlar", "Giderler", "Denetim"}
        assert len(wb.sheetnames) == 4

    def test_export_audit_admin_ok(self, client):
        set_role("admin")
        r = client.get(f"{BASE_URL}/api/export/audit")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        assert "Denetim" in wb.sheetnames

    def test_export_audit_muhasebe_forbidden(self, client):
        set_role("muhasebe")
        r = client.get(f"{BASE_URL}/api/export/audit")
        set_role("admin")
        assert r.status_code == 403

    def test_export_audit_ogretmen_forbidden(self, client):
        set_role("ogretmen")
        r = client.get(f"{BASE_URL}/api/export/audit")
        set_role("admin")
        assert r.status_code == 403

    def test_unknown_kind_404(self, client):
        r = client.get(f"{BASE_URL}/api/export/unknown")
        assert r.status_code == 404


# ------- Reminders -------
class TestReminders:
    def test_reminders_shape(self, client):
        r = client.get(f"{BASE_URL}/api/reminders?days=365")
        assert r.status_code == 200
        items = r.json()
        assert isinstance(items, list)
        if items:
            it = items[0]
            for key in ("student_id", "student_name", "parent_phone",
                        "installment_no", "amount_due", "due_date", "status", "days_diff"):
                assert key in it, f"missing key {key}"
            assert it["status"] in ("overdue", "upcoming")


# ------- Student detail -------
class TestStudentDetail:
    def test_detail_ok(self, client):
        r = client.get(f"{BASE_URL}/api/students/student-001/detail")
        assert r.status_code == 200
        data = r.json()
        assert "student" in data and "payments" in data and "installment_plan" in data
        s = data["student"]
        plan = data["installment_plan"]
        assert len(plan) == s["installment_count"]
        for p in plan:
            for k in ("no", "due_date", "amount", "paid", "remaining", "status"):
                assert k in p
            assert p["status"] in ("paid", "partial", "pending")

    def test_detail_404(self, client):
        r = client.get(f"{BASE_URL}/api/students/nonexistent-xyz/detail")
        assert r.status_code == 404


# ------- Audit log entries created for report + export -------
class TestAuditEntries:
    def test_report_and_export_create_audit(self, client):
        set_role("admin")
        # trigger events
        client.get(f"{BASE_URL}/api/reports/monthly?month=2026-03")
        client.get(f"{BASE_URL}/api/export/students")
        # fetch audit
        r = client.get(f"{BASE_URL}/api/audit-logs?limit=50")
        assert r.status_code == 200
        actions = [a["action"] for a in r.json()]
        assert "report.monthly" in actions
        assert "export.excel" in actions
